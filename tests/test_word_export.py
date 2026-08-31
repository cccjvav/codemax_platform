"""S2-01-4 测试：DDL → Word 数据字典导出（python-docx）。

断言全部落在**导出的字节**上：是不是合法 zip、能不能被 python-docx 读回、
表数/行数/单元格内容与后端解析结果是否一致 —— 不是只测状态码。
"""
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document

from app.tools.sql_ddl import parse_ddl
from app.tools.word import FILENAME, MIME_DOCX, build_data_dictionary

FULL_INIT_SQL = Path(__file__).resolve().parents[1] / "database init" / "full_init.sql"
TABLE_NAMES = ["sys_user", "sys_order", "sys_config", "oauth_client", "oauth_code"]


def _graph() -> dict:
    return parse_ddl(FULL_INIT_SQL.read_text(encoding="utf-8"))


async def _export(client, ddl: str):
    return await client.post("/tools/word-export", json={"ddl": ddl})


# ---------- 纯函数 ----------


def test_build_data_dictionary_produces_readable_docx():
    raw = build_data_dictionary(_graph())
    assert zipfile.is_zipfile(BytesIO(raw))  # .docx 本质是 zip 包
    doc = Document(BytesIO(raw))  # 能被 python-docx 读回才算合法文档
    assert len(doc.tables) == 5


def test_build_data_dictionary_table_contents():
    graph = _graph()
    doc = Document(BytesIO(build_data_dictionary(graph)))

    # 每张 DB 表对应一张 Word 表格，行数 = 表头 + 字段数
    assert [len(t.rows) - 1 for t in doc.tables] == [len(t["columns"]) for t in graph["tables"]]
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["字段", "类型", "主键", "可空", "默认值", "注释"]

    # 抽查 sys_user：id 是主键、username 非空、status 有默认值
    user = {r.cells[0].text: [c.text for c in r.cells] for r in doc.tables[0].rows[1:]}
    assert user["id"][1:4] == ["SERIAL", "✔", "否"]
    assert user["username"][1:4] == ["VARCHAR(50)", "", "否"]
    assert user["nickname"][1:4] == ["VARCHAR(50)", "", "是"]
    assert user["status"][4] == "1"

    # 表名进了标题，外键进了正文
    headings = [p.text for p in doc.paragraphs]
    assert all(name in headings for name in TABLE_NAMES)
    assert "外键关系" in headings
    body = "\n".join(headings)
    assert "sys_order.user_id → sys_user.id" in body
    assert "oauth_code.client_id → oauth_client.id" in body
    assert "共 5 张表，3 条外键关系。" in body


# ---------- 接口 ----------


async def test_word_export_returns_attachment(client):
    r = await _export(client, FULL_INIT_SQL.read_text(encoding="utf-8"))
    assert r.status_code == 200
    assert r.headers["content-type"] == MIME_DOCX
    assert f'attachment; filename="{FILENAME}"' in r.headers["content-disposition"]
    assert zipfile.is_zipfile(BytesIO(r.content))
    assert len(Document(BytesIO(r.content)).tables) == 5


async def test_word_export_rejects_ddl_without_table(client):
    r = await _export(client, "SELECT 1")
    assert r.status_code == 400
    assert "CREATE TABLE" in r.json()["detail"]


async def test_word_export_rejects_empty_ddl(client):
    assert (await _export(client, "")).status_code == 422
