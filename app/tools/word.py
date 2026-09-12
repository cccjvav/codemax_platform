"""ER 解析结果 → Word 数据字典（S2-01-4）。

用 python-docx（不是路线图原先写的 Apache POI —— 那是 Java 库）。
纯函数：吃 parse_ddl 的 {tables, edges}，吐 .docx 字节，不碰 HTTP。
"""
from __future__ import annotations

import re
from io import BytesIO

from docx import Document

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FILENAME = "data_dictionary.docx"

_HEADERS = ("字段", "类型", "主键", "可空", "默认值", "注释")


def build_data_dictionary(graph: dict) -> bytes:
    """把 {tables, edges} 写成 Word 数据字典：每张表一节 + 一张字段表格，末尾列外键。"""
    def safe(value):
        if isinstance(value, str):
            return re.sub(r"[^\x09\x0a\x0d\x20-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]", "", value)
        if isinstance(value, dict):
            return {k: safe(v) for k, v in value.items()}
        if isinstance(value, list):
            return [safe(v) for v in value]
        return value

    graph = safe(graph)
    doc = Document()
    doc.add_heading("数据库数据字典", level=0)
    doc.add_paragraph(f"共 {len(graph['tables'])} 张表，{len(graph['edges'])} 条外键关系。")

    for table in graph["tables"]:
        suffix = f"（{table['comment']}）" if table.get("comment") else ""
        doc.add_heading(f"{table['name']}{suffix}", level=1)
        grid = doc.add_table(rows=1, cols=len(_HEADERS))
        grid.style = "Table Grid"
        for i, head in enumerate(_HEADERS):
            grid.rows[0].cells[i].text = head
            grid.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        for col in table["columns"]:
            cells = grid.add_row().cells
            cells[0].text = col["name"]
            cells[1].text = col["type"]
            cells[2].text = "✔" if col["primary_key"] else ""
            cells[3].text = "否" if not col["nullable"] else "是"
            cells[4].text = "" if col["default"] is None else str(col["default"])
            cells[5].text = col["comment"] or ""

    if graph["edges"]:
        doc.add_heading("外键关系", level=1)
        for e in graph["edges"]:
            doc.add_paragraph(
                f"{e['from_table']}.{e['from_column']} → {e['to_table']}.{e['to_column']}",
                style="List Bullet",
            )

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
